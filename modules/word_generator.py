from docx import Document
import os
import pandas as pd

def update_word_form(template_path, data, save_dir):
    """
    Gera documentos Word individuais para cada linha do DataFrame.

    Args:
        template_path (str): Caminho para o formulário Word existente.
        data (pd.DataFrame): DataFrame contendo os dados.
        save_dir (str): Diretório para salvar os documentos Word gerados.
    """
    # Validação inicial
    if not os.path.exists(template_path):
        print(f"Erro: Template não encontrado em {template_path}")
        return

    if data.empty:
        print("Nenhum dado disponível para atualizar o formulário.")
        return

    if not os.path.exists(save_dir):
        os.makedirs(save_dir)  # Cria o diretório de saída, se não existir

    # Processar cada linha do DataFrame
    for index, row in data.iterrows():
        try:
            # Abrir o template do Word
            doc = Document(template_path)

            # Criar um dicionário para mapear colunas do Excel aos placeholders
            field_mapping = {column: row[column] if not pd.isnull(row[column]) else '' for column in data.columns}

            # Substituir os placeholders no documento Word
            for paragraph in doc.paragraphs:
                for key, value in field_mapping.items():
                    placeholder = f"{{{{{key}}}}}"  # Marcador no formato {{campo}}
                    if placeholder in paragraph.text:
                        paragraph.text = paragraph.text.replace(placeholder, str(value))

            # Definir o nome do arquivo usando "Nome Completo" e "CPF"
            nome_completo = field_mapping.get("Nome Completo", "Desconhecido")
            cpf = field_mapping.get("CPF", "000.000.000-00")
            primeiro_nome = nome_completo.split()[0] if nome_completo else "Desconhecido"

            file_name = f"{primeiro_nome}_{cpf}.docx"
            save_path = os.path.join(save_dir, file_name)

            # Salvar o documento gerado
            doc.save(save_path)
            print(f"Documento salvo: {save_path}")
        
        except Exception as e:
            print(f"Erro ao gerar documento para linha {index + 1}: {str(e)}")
